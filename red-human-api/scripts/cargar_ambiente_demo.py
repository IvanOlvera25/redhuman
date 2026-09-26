"""Carga del AMBIENTE DE DEMOSTRACIÓN desde `Red_Human_AI_Ambiente_Demo.xlsx` — 2026-09-26.

Lee las hojas «Configuración», «Vacantes», «Candidatos», «Contactos demo» y «Usuarios definidos» y deja
en la base dos Cuentas demo («Reclutadora del centro» y «Agencia de talento») con sus Clientes, contactos,
20 vacantes y 514 candidatos repartidos por etapa, con expedientes coherentes a su etapa.

Seguridad (reglas del encargo):
  * Sin `--ejecutar-demo` es un SIMULACRO: hace toda la carga y todas las validaciones dentro de una
    transacción y al final la DESHACE (rollback). Nada queda escrito.
  * Con `--ejecutar-demo` solo confirma (commit) si TODAS las validaciones pasan; si una falla, rollback.
  * Cero comunicaciones: antes de tocar la base se reemplazan por un bloqueo todas las funciones de envío
    (WhatsApp, correo, notificaciones, Teams). El script nunca las llama; si algo lo intentara, el intento
    se cuenta y la carga falla. Además los datos se arman para que ningún job automático tenga a quién
    escribirle: candidatos sin teléfono y con correo `@demo.invalid`, entrevistas humanas en el pasado con
    `recordatorio_enviado_en` puesto, expedientes sin `documentos_hasta`, sin videollamadas agendadas y
    Entrevistas Red Human ya cerradas.
  * Nada se publica: las vacantes quedan «En revisión» (no salen en el portal global ni en el menú del
    WhatsApp compartido) y sin plataformas.
  * Idempotente: los IDs del Excel (VAC-DEMO-###, CAND-DEMO-####) son los `codigo` de vacantes y personas;
    las Cuentas demo se reconocen por su slug `demo-…`. Lo existente se ACTUALIZA, nunca se duplica. Si un
    código demo o el nombre de una Cuenta demo ya pertenecen a datos reales, el script aborta sin escribir.
  * Datos reales intactos: se toma una huella de todas las filas que no son demo antes y después; si
    cambia, rollback. Los 9 administradores se buscan por nombre entre los usuarios YA registrados (nunca
    se crean ni se les cambia correo/contraseña); se vinculan a «Reclutadora del centro» y NO a «Agencia de
    talento». Sus accesos a Cuentas reales no se tocan. Cualquier resultado generado lleva «Simulado».

Uso (desde red-human-api/, con el .env del ambiente cargado; con la API ya arrancada al menos una vez
con esta versión para que el esquema esté al día):
    .venv/Scripts/python.exe scripts/cargar_ambiente_demo.py                   # simulacro (no escribe)
    .venv/Scripts/python.exe scripts/cargar_ambiente_demo.py --ejecutar-demo   # aplica la carga
Opciones:
    --excel RUTA              otro archivo (default: Red_Human_AI_Ambiente_Demo.xlsx en la raíz del repo)
    --correo "Nombre=correo"  (repetible) fija el usuario de un integrante cuando su nombre es ambiguo
"""

import argparse
import hashlib
import inspect as pyinspect
import json
import re
import secrets
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
EXCEL_DEFAULT = RAIZ.parent / "Red_Human_AI_Ambiente_Demo.xlsx"

try:  # la consola de Windows (cp1252) no imprime ✅/❌
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:  # noqa: BLE001
    pass

import openpyxl  # noqa: E402
from sqlalchemy import func, inspect as sa_inspect, or_, select  # noqa: E402

from app.database import Base, SessionLocal, engine, get_db  # noqa: E402
from app.deps import usuario_actual  # noqa: E402
from app.main import app  # noqa: E402  (carga todos los routers: el bloqueo de envíos alcanza a todos)
from app.models import (  # noqa: E402
    DOCUMENTOS_BASE, ETAPAS_CANDIDATO, Candidato, Cliente, ClienteContacto, Cuenta, Documento, Entrevista,
    EntrevistaHumana, Expediente, Postulacion, Usuario, UsuarioCuenta, Vacante, registrar, slugificar,
    texto_sueldo, texto_ubicacion,
)
from app.routers.candidatos import crear_postulacion  # noqa: E402
from app.services import ia  # noqa: E402

ACTOR = "script:cargar_ambiente_demo"
SIMULADO = "Simulado"
PREFIJO_VAC = "VAC-DEMO-"
PREFIJO_CAND = "CAND-DEMO-"
DOMINIO_CORREO = "demo.invalid"  # dominio reservado (RFC 2606): ningún correo puede entregarse
ESTADO_VACANTE = "En revisión"   # decisión del usuario: nada se publica

ETAPA_EXCEL = {  # hoja Candidatos y columnas de la hoja Vacantes → valor interno de la etapa
    "prefiltro": "Prefiltro",
    "entrevista red human": "Entrevista IA",
    "entrevista rh": "Entrevista IA",
    "evaluacion": "Evaluación",
    "entrevista humana": "Entrevista Humana",
    "contratacion": "Contratación",
    "onboarding": "Onboarding",
}
COLUMNAS_ETAPA = ["Prefiltro", "Entrevista RH", "Evaluación", "Entrevista humana", "Contratación", "Onboarding"]
RESULTADO_EXCEL = {"pendiente": ("pendiente", None), "apto": ("cumple", True), "en revision": ("revision", None)}
MODALIDAD = {"presencial": "Presencial", "hibrida": "Híbrido", "hibrido": "Híbrido", "remota": "Remoto", "remoto": "Remoto"}
SENIORITY = {
    "operativo": "Junior", "tecnico": "Semi-senior", "analista": "Semi-senior", "ejecutivo": "Semi-senior",
    "especialista": "Senior", "coordinacion": "Jefatura", "gerencial": "Jefatura", "directivo": "Dirección",
}
ORDEN = {e: i for i, e in enumerate(ETAPAS_CANDIDATO)}


class AbortarCarga(Exception):
    pass


def norm(texto) -> str:
    plano = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", plano).strip().lower()


def slug_cuenta(nombre: str) -> str:
    return "demo-" + slugificar(nombre)


# ============================================================
# Cero comunicaciones
# ============================================================

INTENTOS_BLOQUEADOS: list = []


def bloquear_comunicaciones() -> int:
    """Sustituye TODA función de envío por un bloqueo, en su módulo y en cada módulo `app.*` que la haya
    importado por nombre (`from ..services.whatsapp import enviar_mensaje`)."""
    from app.services import correo, notificaciones, teams, whatsapp

    prefijos = {
        whatsapp: ("enviar", "descargar"),
        correo: ("enviar",),
        notificaciones: ("disparar", "notificar", "enviar"),
        teams: ("crear", "modificar", "actualizar", "cancelar", "enviar"),
    }
    reemplazos = {}
    for modulo, pref in prefijos.items():
        for nombre, obj in list(vars(modulo).items()):
            if not (callable(obj) and nombre.startswith(pref) and getattr(obj, "__module__", "") == modulo.__name__):
                continue
            etiqueta = f"{modulo.__name__}.{nombre}"
            if pyinspect.iscoroutinefunction(obj):
                async def bloqueo(*_a, _e=etiqueta, **_k):
                    INTENTOS_BLOQUEADOS.append(_e)
                    raise RuntimeError(f"Comunicación bloqueada en la carga demo: {_e}")
            else:
                def bloqueo(*_a, _e=etiqueta, **_k):
                    INTENTOS_BLOQUEADOS.append(_e)
                    raise RuntimeError(f"Comunicación bloqueada en la carga demo: {_e}")
            reemplazos[id(obj)] = bloqueo
    for nombre_mod, mod in list(sys.modules.items()):
        if not (nombre_mod == "app" or nombre_mod.startswith("app.")) or mod is None:
            continue
        for attr, obj in list(vars(mod).items()):
            if id(obj) in reemplazos:
                setattr(mod, attr, reemplazos[id(obj)])
    return len(reemplazos)


# ============================================================
# Excel
# ============================================================


def leer_excel(ruta: Path) -> dict:
    if not ruta.exists():
        raise AbortarCarga(f"No encuentro el archivo {ruta}")
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    hojas = {norm(ws.title): ws for ws in wb.worksheets}

    def filas(nombre: str) -> list:
        ws = hojas.get(norm(nombre))
        if ws is None:
            raise AbortarCarga(f"El Excel no tiene la hoja «{nombre}»")
        it = ws.iter_rows(values_only=True)
        encabezados = [norm(h) for h in next(it)]
        salida = []
        for r in it:
            if not r or all(v in (None, "") for v in r):
                continue
            salida.append({encabezados[i]: (v.strip() if isinstance(v, str) else v) for i, v in enumerate(r) if i < len(encabezados)})
        return salida

    ws_conf = hojas.get(norm("Configuración"))
    if ws_conf is None:
        raise AbortarCarga("El Excel no tiene la hoja «Configuración»")
    conf = {norm(r[0]): r[1] for r in ws_conf.iter_rows(values_only=True) if r and r[0]}
    return {
        "config": conf,
        "vacantes": filas("Vacantes"),
        "candidatos": filas("Candidatos"),
        "contactos": filas("Contactos demo"),
        "usuarios": filas("Usuarios definidos"),
    }


def validar_excel(x: dict) -> dict:
    """Consistencia interna del archivo ANTES de tocar la base. Regresa los esperados por vacante/etapa."""
    errores = []
    vacs = {v["id vacante"]: v for v in x["vacantes"]}
    cands = x["candidatos"]
    esperados_conf = {k: int(x["config"].get(k) or 0) for k in ("cuentas", "vacantes", "candidatos")}
    if len(vacs) != len(x["vacantes"]):
        errores.append("IDs de vacante repetidos")
    if len({c["id candidato"] for c in cands}) != len(cands):
        errores.append("IDs de candidato repetidos")
    if esperados_conf["vacantes"] and len(vacs) != esperados_conf["vacantes"]:
        errores.append(f"Configuración dice {esperados_conf['vacantes']} vacantes y la hoja trae {len(vacs)}")
    if esperados_conf["candidatos"] and len(cands) != esperados_conf["candidatos"]:
        errores.append(f"Configuración dice {esperados_conf['candidatos']} candidatos y la hoja trae {len(cands)}")
    cuentas = {v["cuenta"] for v in x["vacantes"]}
    if esperados_conf["cuentas"] and len(cuentas) != esperados_conf["cuentas"]:
        errores.append(f"Configuración dice {esperados_conf['cuentas']} cuentas y las vacantes usan {len(cuentas)}")
    for vid in vacs:
        if not str(vid).startswith(PREFIJO_VAC):
            errores.append(f"{vid}: una vacante demo debe empezar con {PREFIJO_VAC}")

    esperado = defaultdict(Counter)  # vacante → etapa → n
    for c in cands:
        if not str(c["id candidato"]).startswith(PREFIJO_CAND):
            errores.append(f"{c['id candidato']}: un candidato demo debe empezar con {PREFIJO_CAND}")
        v = vacs.get(c["id vacante"])
        if not v:
            errores.append(f"{c['id candidato']}: vacante {c['id vacante']} inexistente")
            continue
        if (c["cuenta"], c["cliente"], c["puesto"]) != (v["cuenta"], v["cliente"], v["puesto"]):
            errores.append(f"{c['id candidato']}: cuenta/cliente/puesto no coinciden con {c['id vacante']}")
        etapa = ETAPA_EXCEL.get(norm(c["etapa"]))
        if not etapa:
            errores.append(f"{c['id candidato']}: etapa desconocida «{c['etapa']}»")
            continue
        if norm(c.get("resultado de ejemplo")) not in RESULTADO_EXCEL:
            errores.append(f"{c['id candidato']}: resultado desconocido «{c.get('resultado de ejemplo')}»")
        esperado[c["id vacante"]][etapa] += 1
    for vid, v in vacs.items():
        declarado = {ETAPA_EXCEL[norm(col)]: int(v.get(norm(col)) or 0) for col in COLUMNAS_ETAPA}
        if sum(declarado.values()) != int(v.get("candidatos") or 0):
            errores.append(f"{vid}: la suma por etapa no da el total de candidatos")
        if {e: n for e, n in declarado.items() if n} != dict(esperado[vid]):
            errores.append(f"{vid}: la hoja Candidatos no coincide con los conteos por etapa de la hoja Vacantes")
    if errores:
        raise AbortarCarga("El Excel no es consistente:\n  - " + "\n  - ".join(errores[:30]))
    return {vid: dict(c) for vid, c in esperado.items()}


# ============================================================
# Utilidades de escritura idempotente
# ============================================================


def _comparable(v):
    if isinstance(v, datetime):
        return (v.astimezone(timezone.utc).replace(tzinfo=None) if v.tzinfo else v).isoformat()
    if isinstance(v, (dict, list)):
        return json.dumps(v, sort_keys=True, ensure_ascii=False, default=str)
    return v


def poner(obj, **campos) -> bool:
    """Asigna SOLO lo que cambió: así una segunda pasada no emite UPDATE (ni dispara `onupdate`)."""
    cambio = False
    for k, v in campos.items():
        if _comparable(getattr(obj, k)) != _comparable(v):
            setattr(obj, k, v)
            cambio = True
    return cambio


class Marcador:
    def __init__(self):
        self.creados = Counter()
        self.actualizados = Counter()
        self.avisos: list = []

    def nuevo(self, tipo: str):
        self.creados[tipo] += 1

    def cambio(self, tipo: str, hubo: bool):
        if hubo:
            self.actualizados[tipo] += 1


def _num(codigo: str) -> int:
    m = re.search(r"(\d+)$", codigo)
    return int(m.group(1)) if m else 0


# ============================================================
# Cuentas, usuarios, clientes y contactos
# ============================================================


def cuenta_demo(db, nombre: str, mk: Marcador) -> Cuenta:
    slug = slug_cuenta(nombre)
    c = db.query(Cuenta).filter(Cuenta.slug == slug).first()
    homonima = (
        db.query(Cuenta)
        .filter(func.lower(Cuenta.nombre) == nombre.lower(), or_(Cuenta.slug.is_(None), Cuenta.slug != slug))
        .first()
    )
    if homonima:
        raise AbortarCarga(
            f"Ya existe una Cuenta REAL llamada «{homonima.nombre}» (id {homonima.id}, slug «{homonima.slug}»). "
            "La carga demo no la toca: renómbrala o usa otra base."
        )
    if c is None:
        c = Cuenta(nombre=nombre, slug=slug)
        db.add(c)
        mk.nuevo("cuenta")
    hubo = poner(
        c, nombre=nombre, nombre_comercial=nombre, razon_social=f"{nombre} S.A. de C.V. (DEMO)",
        contacto_nombre="Ambiente demo (datos ficticios)", correo_comunicacion="", whatsapp_comunicacion="",
        whatsapp_exclusivo=False, estado="Activa", eliminada_en=None, eliminada_por="",
    )
    mk.cambio("cuenta", hubo and c.id is not None)
    db.flush()
    return c


def buscar_usuario(db, nombre: str, overrides: dict):
    """Un usuario YA registrado cuyo nombre (o correo) contiene todas las palabras del nombre del Excel.
    Nunca crea ni inventa: si no hay uno único, regresa (None, motivo)."""
    if norm(nombre) in overrides:
        u = db.query(Usuario).filter(func.lower(Usuario.correo) == overrides[norm(nombre)].lower()).first()
        return (u, None) if u else (None, f"--correo indica {overrides[norm(nombre)]}, que no está registrado")
    palabras = norm(nombre).split()
    hallados = []
    for u in db.query(Usuario).filter(Usuario.activo.is_(True)).order_by(Usuario.id).all():
        tokens = set(norm(u.nombre).split()) | set(re.split(r"[._\-@+]", norm(u.correo)))
        if all(p in tokens for p in palabras):
            hallados.append(u)
    if len(hallados) == 1:
        return hallados[0], None
    if not hallados:
        return None, "no hay un usuario registrado con ese nombre"
    return None, "nombre ambiguo: " + ", ".join(f"{u.nombre} <{u.correo}>" for u in hallados) + " (usa --correo)"


def vincular_admins(db, x: dict, cuentas: dict, overrides: dict, mk: Marcador) -> dict:
    """Administradores de «Reclutadora del centro»: vinculados a ella y NO a «Agencia de talento»
    (decisión del usuario: la exclusividad es entre las Cuentas demo; sus Cuentas reales no se tocan)."""
    encontrados = {}
    ids_demo = {c.id for c in cuentas.values()}
    for fila in x["usuarios"]:
        cuenta_nombre, nombre, rol = fila.get("cuenta"), fila.get("usuario solicitado"), fila.get("rol")
        if norm(rol) != "administrador":
            mk.avisos.append(f"Usuarios definidos · {cuenta_nombre} · «{nombre}»: {rol} — se omite (sin datos que cargar).")
            continue
        cuenta = cuentas.get(norm(cuenta_nombre))
        if cuenta is None:
            mk.avisos.append(f"Usuarios definidos · «{nombre}»: la Cuenta «{cuenta_nombre}» no existe en el Excel.")
            continue
        u, motivo = buscar_usuario(db, nombre, overrides)
        if u is None:
            mk.avisos.append(f"Administrador «{nombre}»: {motivo}. No se crea ni se inventa (queda pendiente).")
            continue
        encontrados[nombre] = u
        if u.rol != "Administrador":
            mk.avisos.append(f"Administrador «{nombre}» ({u.correo}): rol {u.rol} → Administrador.")
            u.rol = "Administrador"
        vinculos = {uc.cuenta_id: uc for uc in db.query(UsuarioCuenta).filter(UsuarioCuenta.usuario_id == u.id).all()}
        if cuenta.id not in vinculos:
            db.add(UsuarioCuenta(usuario_id=u.id, cuenta_id=cuenta.id))
            mk.nuevo("vinculo_usuario_cuenta")
        for cid, uc in vinculos.items():
            if cid in ids_demo and cid != cuenta.id:  # exclusividad dentro del ambiente demo
                db.delete(uc)
                mk.avisos.append(f"Administrador «{nombre}»: se retira su acceso a la otra Cuenta demo (id {cid}).")
    db.flush()
    return encontrados


def cargar_clientes(db, x: dict, cuentas: dict, admins: dict, mk: Marcador) -> dict:
    clientes = {}
    for v in x["vacantes"]:
        cuenta = cuentas[norm(v["cuenta"])]
        clave = (cuenta.id, norm(v["cliente"]))
        if clave in clientes:
            continue
        cl = db.query(Cliente).filter(Cliente.cuenta_id == cuenta.id, func.lower(Cliente.nombre) == v["cliente"].lower()).first()
        if cl is None:
            cl = Cliente(cuenta_id=cuenta.id, nombre=v["cliente"])
            db.add(cl)
            mk.nuevo("cliente")
            nuevo = True
        else:
            nuevo = False
        hubo = poner(cl, razon_social=f"{v['cliente']} S.A. de C.V. (DEMO)", nombre_comercial=v["cliente"], estado="Activo")
        mk.cambio("cliente", hubo and not nuevo)
        clientes[clave] = cl
    db.flush()

    # Contactos: la MISMA persona registrada (nombre, correo y WhatsApp de su usuario, sin cambiarle nada);
    # lo que la distingue es la Cuenta, el Cliente y la función.
    for fila in x["contactos"]:
        cuenta = cuentas.get(norm(fila["cuenta"]))
        cl = clientes.get((cuenta.id, norm(fila["cliente"]))) if cuenta else None
        if cl is None:
            mk.avisos.append(f"Contacto «{fila['nombre del integrante']}»: el Cliente «{fila['cliente']}» no tiene vacantes en el Excel.")
            continue
        integrante = fila["nombre del integrante"]
        u = admins.get(integrante)
        nombre, correo, telefono = (u.nombre, u.correo, u.telefono or "") if u else (integrante, "", "")
        if u is None:
            mk.avisos.append(f"Contacto «{integrante}» en {cl.nombre}: sin usuario registrado, queda sin correo ni WhatsApp.")
        existentes = db.query(ClienteContacto).filter(ClienteContacto.cliente_id == cl.id).all()
        cc = next((c for c in existentes if correo and (c.correo or "").lower() == correo.lower()), None) or next(
            (c for c in existentes if norm(c.nombre) in (norm(nombre), norm(integrante))), None
        )
        if cc is None:
            cc = ClienteContacto(cliente_id=cl.id, nombre=nombre)
            db.add(cc)
            mk.nuevo("contacto")
            nuevo = True
        else:
            nuevo = False
        hubo = poner(cc, nombre=nombre, apellidos="", puesto=fila.get("puesto sugerido") or "", correo=correo, telefono=telefono)
        mk.cambio("contacto", hubo and not nuevo)
    db.flush()
    return clientes


# ============================================================
# Vacantes
# ============================================================


def preguntas_cerradas(requisitos: list) -> list:
    return [
        {"pregunta": f"¿Cuentas con {r[0].lower() + r[1:]}?", "tipo": "si_no", "valida": r, "respuesta_esperada": "Sí",
         "descarta": True, "opciones": ["Sí", "No", "Parcial"]}
        for r in requisitos
    ]


def cargar_vacantes(db, x: dict, cuentas: dict, clientes: dict, admins: dict, base: datetime, mk: Marcador) -> dict:
    vacantes = {}
    responsables = sorted(admins.values(), key=lambda u: u.id)
    for v in x["vacantes"]:
        codigo = v["id vacante"]
        cuenta = cuentas[norm(v["cuenta"])]
        cl = clientes[(cuenta.id, norm(v["cliente"]))]
        vac = db.query(Vacante).filter(Vacante.codigo == codigo).first()
        if vac is not None and vac.cuenta_id != cuenta.id:
            raise AbortarCarga(f"El código {codigo} ya existe en otra Cuenta (id {vac.cuenta_id}); no se toca.")
        nuevo = vac is None
        if nuevo:
            vac = Vacante(codigo=codigo, titulo=v["puesto"])
            db.add(vac)
            mk.nuevo("vacante")
        reqs = [r.strip() for r in str(v.get("requisitos indispensables") or "").split(";") if r.strip()]
        desde, hasta = int(v["sueldo desde"]), int(v["sueldo hasta"])
        periodicidad = norm(v.get("periodicidad")) or "mensual"
        n = _num(codigo)
        responsable = responsables[n % len(responsables)] if responsables and norm(v["cuenta"]) == "reclutadora del centro" else None
        hubo = poner(
            vac, titulo=v["puesto"], slug=f"{slugificar(v['puesto'])}-{codigo.lower()}", area=v.get("area") or "",
            empresa=v["cliente"], cuenta_id=cuenta.id, cliente_id=cl.id, mostrar_cliente_candidato=True,
            responsable_id=responsable.id if responsable else None, colaboradores_ids=[],
            ubicacion_estado=v.get("estado") or "", ubicacion_municipio=v.get("municipio") or "",
            ubicacion=texto_ubicacion(v.get("estado") or "", v.get("municipio") or ""),
            modalidad=MODALIDAD.get(norm(v.get("modalidad")), "Presencial"),
            sueldo_desde=desde, sueldo_hasta=hasta, sueldo_moneda=v.get("moneda") or "MXN", sueldo_periodicidad=periodicidad,
            sueldo=texto_sueldo(desde, hasta, v.get("moneda") or "MXN", periodicidad),
            seniority=SENIORITY.get(norm(v.get("nivel")), ""), descripcion=v.get("descripcion breve") or "",
            requisitos=" · ".join(reqs), preguntas_filtro=preguntas_cerradas(reqs), preguntas_filtro_whatsapp=[],
            enfoque_entrevista="profesional", estado=ESTADO_VACANTE, eliminada_en=None, eliminada_por="",
            plataformas=[], publicaciones={}, publicada_en=None,
            avisos_cumplimiento=[f"Vacante {v.get('marca demo') or 'FICTICIA / DEMO'}: ambiente de demostración, no publicar."],
            creada_en=base - timedelta(days=45 - n),
        )
        mk.cambio("vacante", hubo and not nuevo)
        vacantes[codigo] = vac
    db.flush()
    return vacantes


# ============================================================
# Candidatos y su expediente coherente por etapa
# ============================================================


def _evaluacion(titulo: str, reqs: list, score: int, recomendacion: str) -> dict:
    return {
        "simulado": True,
        "resumen": f"{SIMULADO} — evaluación de ejemplo para la demostración ({titulo}); no es una entrevista real del agente.",
        "fortalezas": [f"{SIMULADO}: experiencia declarada en {reqs[0].lower()}" if reqs else f"{SIMULADO}: experiencia afín al puesto",
                       f"{SIMULADO}: comunicación clara y ordenada"],
        "riesgos": [f"{SIMULADO}: validar {reqs[-1].lower()} con evidencia" if reqs else f"{SIMULADO}: validar referencias"],
        "areas_desarrollo": [],
        "calif_experiencia": round(score / 10, 1),
        "calif_comunicacion": round(min(10, score / 10 + 0.5), 1),
        "match_perfil": score,
        "recomendacion": recomendacion,
        "evidencia": f"{SIMULADO} — respuestas de ejemplo generadas para el ambiente demo.",
        "faltante": [],
    }


def _transcript(titulo: str, reqs: list) -> list:
    tema = reqs[0].lower() if reqs else "tu experiencia"
    return [
        {"rol": "assistant", "texto": ia.mensaje_inicial_entrevista(titulo)},
        {"rol": "user", "texto": f"[{SIMULADO}] Sí, claro, comencemos."},
        {"rol": "assistant", "texto": f"[{SIMULADO}] Cuéntame de tu experiencia con {tema}."},
        {"rol": "user", "texto": f"[{SIMULADO}] Llevo varios años trabajando en {tema}; respuesta de ejemplo para la demo."},
        {"rol": "assistant", "texto": f"[{SIMULADO}] ¿Qué resultado concreto destacarías de ese trabajo?"},
        {"rol": "user", "texto": f"[{SIMULADO}] Mejoré los tiempos de atención de mi equipo; dato ficticio."},
        {"rol": "assistant", "texto": f"{ia.DESPEDIDA_ENTREVISTA}. [{SIMULADO}]"},
    ]


def _entrevista_ia(db, c, p, vac, reqs, score, rec, t, mk) -> None:
    e = next((e for e in p.entrevistas if (e.guion or {}).get("simulado")), None)
    nuevo = e is None
    if nuevo:
        e = Entrevista(codigo="TMP", candidato_id=c.id, postulacion_id=p.id, token=secrets.token_urlsafe(24))
        db.add(e)
        db.flush()
        e.codigo = f"ENT-{300 + e.id}"
        p.entrevistas.append(e)
        mk.nuevo("entrevista_red_human")
    hubo = poner(
        e, tipo="texto", estado="evaluada", cierre="texto", motivo="",
        guion={"enfoque": vac.enfoque_entrevista or "profesional", "temas": reqs, "simulado": True},
        transcript=_transcript(vac.titulo, reqs), evaluacion=_evaluacion(vac.titulo, reqs, score, rec),
        consentimiento=True, consentimiento_fecha=t, programada_para=t, iniciada_en=t,
        finalizada_en=t + timedelta(minutes=18), ultima_actividad_en=t + timedelta(minutes=18), liga_meet="",
    )
    mk.cambio("entrevista_red_human", hubo and not nuevo)


def _entrevista_humana(db, c, p, vac, etapa, apto, t, contactos, n, mk) -> None:
    eh = p.entrevistas_humanas[0] if p.entrevistas_humanas else None
    nuevo = eh is None
    if nuevo:
        eh = EntrevistaHumana(candidato_id=c.id, postulacion_id=p.id, token=secrets.token_urlsafe(24))
        db.add(eh)
        p.entrevistas_humanas.append(eh)
        mk.nuevo("entrevista_humana")
    realizada = not (etapa == "Entrevista Humana" and not apto)
    if contactos:
        cc = contactos[n % len(contactos)]
        quien = dict(tipo="externo", contacto_id=cc.id, entrevistador=cc.nombre, correo_externo=cc.correo or "",
                     whatsapp_externo=cc.telefono or "", usuario_id=None)
    else:
        quien = dict(tipo="externo", contacto_id=None, entrevistador=f"Entrevistador demo ({SIMULADO})",
                     correo_externo="", whatsapp_externo="", usuario_id=None)
    hubo = poner(
        eh, **quien, fecha=t, modalidad="Videollamada", liga="https://demo.invalid/entrevista-simulada",
        ubicacion="", telefono_contacto="", cancelada=False, teams_evento_id="",
        realizada=realizada,
        resultado="aprobado" if realizada else "",
        recomendacion="avanzar" if realizada else "",
        resultado_capturado_por="rh" if realizada else "",
        evaluada_en=t + timedelta(hours=2) if realizada else None,
        comentario=(f"{SIMULADO} — notas de ejemplo: buena disposición y experiencia acorde; decisión de ejemplo, no real."
                    if realizada else f"{SIMULADO} — entrevista de ejemplo pendiente de capturar resultado."),
        # en el pasado y con el recordatorio marcado: el job de recordatorios nunca la toma
        recordatorio_enviado_en=t - timedelta(days=1),
    )
    mk.cambio("entrevista_humana", hubo and not nuevo)


def _expediente(db, c, p, vac, cuenta_rs, etapa, apto, n, t, mk) -> None:
    exp = p.expediente
    nuevo = exp is None
    if nuevo:
        exp = Expediente(candidato_id=c.id, postulacion_id=p.id, token=secrets.token_urlsafe(24))
        db.add(exp)
        p.expediente = exp
        mk.nuevo("expediente")
    onboarding = etapa == "Onboarding"
    medio = round((vac.sueldo_desde + vac.sueldo_hasta) / 2 / 500) * 500
    hubo = poner(
        exp, candidato_id=c.id, puesto=vac.titulo, sueldo=texto_sueldo(medio, medio, vac.sueldo_moneda, vac.sueldo_periodicidad),
        tipo_contratacion="Tiempo indeterminado", ubicacion=vac.ubicacion, jefe_directo=f"Jefatura de {vac.area} ({SIMULADO})",
        fecha_ingreso=t + timedelta(days=3 if onboarding else 14), empresa=cuenta_rs,
        condiciones_guardadas_en=t, duracion_contrato=None, duracion_unidad="", fecha_termino=None,
        contrato="Firmado" if onboarding else "Pendiente", alta_administrativa="Pendiente", equipo_accesos="Pendiente",
        estado="completo" if onboarding else "integracion", seleccionado_por=SIMULADO,
        instrucciones_ingreso=f"{SIMULADO} — instrucciones de ejemplo: presentarse a las 9:00 con identificación.",
        documentos_hasta=None,  # sin recordatorios automáticos
        ultimo_recordatorio_en=None, documentos_vencidos_avisado=False, recordatorios_enviados=0,
    )
    mk.cambio("expediente", hubo and not nuevo)
    db.flush()

    entregados = len(DOCUMENTOS_BASE) if onboarding else ((n % 4) + 2 if apto else n % 2)
    existentes = {d.tipo: d for d in exp.documentos}
    for i, tipo in enumerate(DOCUMENTOS_BASE):
        d = existentes.get(tipo)
        nuevo_doc = d is None
        if nuevo_doc:
            d = Documento(expediente_id=exp.id, tipo=tipo)
            db.add(d)
            exp.documentos.append(d)
            mk.nuevo("documento")
        recibido = i < entregados
        hubo = poner(
            d, obligatorio=True, estado="recibido" if recibido else "pendiente", archivo="", nombre_archivo="", mime="", tamano=0,
            notas_ia=f"{SIMULADO} — documento de ejemplo, no se recibió ningún archivo real." if recibido else "",
            validacion={}, revisado_por=f"{SIMULADO} (carga demo)" if recibido else "",
            subido_en=t + timedelta(days=1) if recibido else None,
            solicitado_en=t, solicitado_canal="simulado",
            solicitudes=[{"en": t.isoformat(), "canal": "simulado", "tipo": "solicitud", "por": SIMULADO}],
            recibido_en=t + timedelta(days=1) if recibido else None, recibido_canal="rh" if recibido else "",
        )
        mk.cambio("documento", hubo and not nuevo_doc)


def cargar_candidatos(db, x: dict, cuentas: dict, vacantes: dict, base: datetime, mk: Marcador) -> None:
    contactos_por_cliente = {}
    for cand in x["candidatos"]:
        codigo = cand["id candidato"]
        vac = vacantes[cand["id vacante"]]
        cuenta = cuentas[norm(cand["cuenta"])]
        etapa = ETAPA_EXCEL[norm(cand["etapa"])]
        estado, apto = RESULTADO_EXCEL[norm(cand.get("resultado de ejemplo"))]
        score = int(cand.get("score de ejemplo") or 0)
        n = _num(codigo)
        nivel = ORDEN[etapa]
        creado = base - timedelta(days=4 + nivel * 6 + n % 5)
        reqs = [r for r in (vac.requisitos or "").split(" · ") if r]

        c = db.query(Candidato).filter(Candidato.codigo == codigo).first()
        if c is not None and c.cuenta_id != cuenta.id:
            raise AbortarCarga(f"El código {codigo} ya existe en otra Cuenta (id {c.cuenta_id}); no se toca.")
        nuevo = c is None
        if nuevo:
            c = Candidato(codigo=codigo, nombre=cand["nombre ficticio"], cuenta_id=cuenta.id)
            db.add(c)
            db.flush()
            mk.nuevo("candidato")
        hubo = poner(
            c, nombre=cand["nombre ficticio"], cuenta_id=cuenta.id, es_prueba=False, eliminado_en=None, eliminado_por="",
            correo=f"{codigo.lower()}@{DOMINIO_CORREO}", telefono="", wa_id="", wa_nombre="",  # sin WhatsApp
            ubicacion=vac.ubicacion, experiencia=f"{2 + n % 9} años en {vac.area.lower()}", fuente="Demo",
            cv_datos={"simulado": True, "resumen_profesional": f"Perfil ficticio ({SIMULADO}) para {vac.titulo}."},
            creado_en=creado,
        )
        mk.cambio("candidato", hubo and not nuevo)

        p = next((p for p in sorted(c.postulaciones, key=lambda q: q.id) if p.vacante_id == vac.id), None)
        nueva_p = p is None
        if nueva_p:
            p = crear_postulacion(db, c, vac, cuenta.id, "rh_directo", etapa=etapa)
            mk.nuevo("postulacion")
        avanzado = nivel >= ORDEN["Entrevista IA"]
        consentimiento = avanzado or n % 3 != 0
        analisis = {"simulado": True, "leyenda": f"{SIMULADO}: datos de ejemplo del ambiente demo."}
        if avanzado:
            analisis.update({
                "origen": "cv", "ia": False,
                "requisitos_cumplidos": [f"{r} ({SIMULADO})" for r in reqs],
                "brechas": [], "fortalezas_cv": [f"{SIMULADO}: trayectoria afín a {vac.titulo.lower()}"],
                "compatibilidad_cv": f"{SIMULADO} — compatibilidad de ejemplo: {score}/100.",
                "experiencia_relevante_cv": f"{SIMULADO} — experiencia de ejemplo en {vac.area.lower()}.",
                "alertas": [], "datos_faltantes": [],
                "prefiltro_resultado": "cumple",
                "respuestas_prefiltro": [
                    {"criterio": r, "pregunta": f"¿Cuentas con {r.lower()}?", "respuesta": f"Sí ({SIMULADO})", "cumple": True}
                    for r in reqs
                ],
            })
        historial = list(p.historial or [])
        if not any(h.get("evento") == "carga_demo" for h in historial):
            historial.append({"evento": "carga_demo", "texto": f"Expediente {SIMULADO}: cargado para el ambiente de demostración; "
                              "ningún resultado corresponde a una entrevista o decisión real.", "usuario": ACTOR,
                              "fecha": creado.isoformat()})
        hubo = poner(
            p, cuenta_id=cuenta.id, vacante_id=vac.id, activa=True, motivo_cierre="", cerrada_en=None, origen="rh_directo",
            es_prueba=False, etapa=etapa, estado=estado, resultado_apto=apto, score=score if avanzado else 0,
            evidencia=f"{SIMULADO} — resultado de ejemplo ({cand.get('resultado de ejemplo')})." if avanzado else "Prefiltro en curso.",
            analisis=analisis, prefiltro_completo=avanzado, consentimiento=consentimiento,
            consentimiento_fecha=creado if consentimiento else None, historial=historial,
            ultima_actividad_en=base - timedelta(days=n % 3), creado_en=creado,
            videollamada_agendada_en=None, videollamada_liga="", videollamada_aviso_noshow_enviado=False,
        )
        mk.cambio("postulacion", hubo and not nueva_p)

        t_ent = creado + timedelta(days=2)
        if avanzado:
            rec = "avanzar" if apto else "revision"
            _entrevista_ia(db, c, p, vac, reqs, score, rec, t_ent, mk)
        if nivel >= ORDEN["Entrevista Humana"]:
            if vac.cliente_id not in contactos_por_cliente:
                contactos_por_cliente[vac.cliente_id] = (
                    db.query(ClienteContacto).filter(ClienteContacto.cliente_id == vac.cliente_id).order_by(ClienteContacto.id).all()
                )
            _entrevista_humana(db, c, p, vac, etapa, apto, t_ent + timedelta(days=4), contactos_por_cliente[vac.cliente_id], n, mk)
        if nivel >= ORDEN["Contratación"]:
            rs = vac.cliente.razon_social if vac.cliente and vac.cliente.razon_social else cuenta.razon_social
            _expediente(db, c, p, vac, rs, etapa, apto, n, t_ent + timedelta(days=8), mk)
        if n % 50 == 0:
            db.flush()
    db.flush()


def cargar(db, x: dict, overrides: dict, base: datetime) -> tuple:
    mk = Marcador()
    nombres = sorted({v["cuenta"] for v in x["vacantes"]})
    cuentas = {norm(n): cuenta_demo(db, n, mk) for n in nombres}
    admins = vincular_admins(db, x, cuentas, overrides, mk)
    clientes = cargar_clientes(db, x, cuentas, admins, mk)
    vacantes = cargar_vacantes(db, x, cuentas, clientes, admins, base, mk)
    cargar_candidatos(db, x, cuentas, vacantes, base, mk)
    return mk, cuentas, admins


# ============================================================
# Huellas: datos reales intactos e idempotencia
# ============================================================


def ids_demo(db) -> dict:
    cuentas = {c for (c,) in db.query(Cuenta.id).filter(Cuenta.slug.like("demo-%")).all()}
    if not cuentas:
        return {"cuentas": set(), "candidatos": set(), "postulaciones": set(), "expedientes": set(), "clientes": set()}
    cands = {i for (i,) in db.query(Candidato.id).filter(Candidato.cuenta_id.in_(cuentas)).all()}
    posts = {i for (i,) in db.query(Postulacion.id).filter(Postulacion.cuenta_id.in_(cuentas)).all()}
    exps = {i for (i,) in db.query(Expediente.id).filter(Expediente.postulacion_id.in_(posts)).all()} if posts else set()
    clis = {i for (i,) in db.query(Cliente.id).filter(Cliente.cuenta_id.in_(cuentas)).all()}
    return {"cuentas": cuentas, "candidatos": cands, "postulaciones": posts, "expedientes": exps, "clientes": clis}


def huella(db, solo_reales: bool) -> dict:
    """sha256 por tabla. `solo_reales`: excluye las filas del ambiente demo (y la columna `rol` de
    usuarios, que la carga puede ajustar a propósito y se reporta aparte)."""
    demo = ids_demo(db) if solo_reales else None
    tablas = set(sa_inspect(db.get_bind()).get_table_names())
    salida = {}
    for tabla in Base.metadata.sorted_tables:
        if tabla.name not in tablas:
            continue
        cols = [c for c in tabla.columns if not (solo_reales and tabla.name == "usuarios" and c.name == "rol")]
        pk = list(tabla.primary_key.columns)
        h = hashlib.sha256()
        n = 0
        for fila in db.execute(select(*cols).order_by(*pk)).mappings():
            if tabla.name == "bitacora" and fila.get("accion") == "carga_ambiente_demo":
                continue
            if demo is not None and (
                (tabla.name == "cuentas" and fila["id"] in demo["cuentas"])
                or (tabla.name == "clientes" and fila["id"] in demo["clientes"])
                or (tabla.name != "bitacora" and fila.get("cuenta_id") in demo["cuentas"])
                or fila.get("candidato_id") in demo["candidatos"]
                or fila.get("postulacion_id") in demo["postulaciones"]
                or fila.get("expediente_id") in demo["expedientes"]
                or (tabla.name != "vacantes" and fila.get("cliente_id") in demo["clientes"])
            ):
                continue
            h.update(json.dumps({k: _comparable(v) for k, v in fila.items()}, sort_keys=True, default=str).encode())
            n += 1
        salida[tabla.name] = (n, h.hexdigest())
    return salida


# ============================================================
# Validación a través de la API (misma sesión, sin escribir)
# ============================================================


class Validador:
    def __init__(self):
        self.ok = 0
        self.fallas: list = []

    def check(self, cond: bool, msg: str):
        if cond:
            self.ok += 1
            print(f"  ✅ {msg}")
        else:
            self.fallas.append(msg)
            print(f"  ❌ {msg}")


def validar(db, x: dict, esperado: dict, cuentas: dict, admins: dict, val: Validador) -> None:
    from fastapi.testclient import TestClient

    vac_excel = {v["id vacante"]: v for v in x["vacantes"]}
    cand_excel = {c["id candidato"]: c for c in x["candidatos"]}
    por_cuenta = defaultdict(set)
    for vid, v in vac_excel.items():
        por_cuenta[norm(v["cuenta"])].add(vid)

    # --- base de datos ---
    print("\n— Conteos en la base")
    vac_db = db.query(Vacante).filter(Vacante.codigo.like(f"{PREFIJO_VAC}%")).all()
    cand_db = db.query(Candidato).filter(Candidato.codigo.like(f"{PREFIJO_CAND}%")).all()
    val.check(len(vac_db) == len(vac_excel) == 20, f"vacantes demo: {len(vac_db)} (Excel {len(vac_excel)}, esperado 20)")
    val.check(len(cand_db) == len(cand_excel) == 514, f"candidatos demo: {len(cand_db)} (Excel {len(cand_excel)}, esperado 514)")
    posts = db.query(Postulacion).join(Candidato, Postulacion.candidato_id == Candidato.id).filter(Candidato.codigo.like(f"{PREFIJO_CAND}%")).all()
    val.check(len(posts) == len(cand_excel), f"una postulación por candidato demo ({len(posts)})")
    reales_en_demo = [p.codigo for p in posts if p.vacante is None or not p.vacante.codigo.startswith(PREFIJO_VAC)]
    val.check(not reales_en_demo, "ninguna postulación demo apunta a una vacante real")
    etapa_db = Counter(p.etapa for p in posts)
    etapa_xl = Counter(ETAPA_EXCEL[norm(c["etapa"])] for c in cand_excel.values())
    val.check(etapa_db == etapa_xl, "etapas en la base = Excel: " + ", ".join(f"{e} {etapa_db.get(e, 0)}" for e in ETAPAS_CANDIDATO))
    cruzados = [p.codigo for p in posts if not (p.cuenta_id == p.vacante.cuenta_id == p.candidato.cuenta_id == (p.vacante.cliente.cuenta_id if p.vacante.cliente else p.cuenta_id))]
    val.check(not cruzados, "postulación, persona, vacante y Cliente siempre en la misma Cuenta")
    val.check(all(v.estado == ESTADO_VACANTE and not v.plataformas for v in vac_db), f"ninguna vacante demo publicada (todas «{ESTADO_VACANTE}», sin plataformas)")
    val.check(all(not c.telefono and not c.wa_id and c.correo.endswith("@" + DOMINIO_CORREO) for c in cand_db),
              f"candidatos sin WhatsApp y con correo no entregable (@{DOMINIO_CORREO})")
    incoherentes = []
    for p in posts:
        nivel = ORDEN[p.etapa]
        ent = [e for e in p.entrevistas if e.estado == "evaluada" and SIMULADO in (e.evaluacion or {}).get("resumen", "")]
        if nivel >= ORDEN["Entrevista IA"] and not ent:
            incoherentes.append(f"{p.codigo} sin Entrevista Red Human simulada")
        if nivel >= ORDEN["Entrevista Humana"] and not p.entrevistas_humanas:
            incoherentes.append(f"{p.codigo} sin Entrevista Humana")
        if nivel >= ORDEN["Contratación"] and not (p.expediente and len(p.expediente.documentos) == len(DOCUMENTOS_BASE)):
            incoherentes.append(f"{p.codigo} sin expediente completo")
        if p.etapa == "Onboarding" and p.expediente and p.expediente.progreso != 100:
            incoherentes.append(f"{p.codigo} en Onboarding con expediente al {p.expediente.progreso}%")
    val.check(not incoherentes, "expedientes coherentes con la etapa (con leyenda «Simulado»)" + (f": {incoherentes[:3]}" if incoherentes else ""))
    ids_d = {c.id for c in cuentas.values()}
    riesgo = []
    for p in posts:
        for eh in p.entrevistas_humanas:
            if eh.fecha and _comparable(eh.fecha) > _comparable(datetime.now(timezone.utc)) and not eh.recordatorio_enviado_en:
                riesgo.append(eh.id)
        if p.expediente and p.expediente.documentos_hasta:
            riesgo.append(p.codigo)
        if p.videollamada_agendada_en:
            riesgo.append(p.codigo)
    val.check(not riesgo, "ningún job automático tiene a quién escribir (recordatorios, no-show, documentos)")
    recl = cuentas.get("reclutadora del centro")
    agencia = cuentas.get("agencia de talento")
    for nombre, u in admins.items():
        vinc = {cid for (cid,) in db.query(UsuarioCuenta.cuenta_id).filter(UsuarioCuenta.usuario_id == u.id).all()}
        val.check(recl.id in vinc and (agencia is None or agencia.id not in vinc) and (vinc & ids_d) == {recl.id},
                  f"administrador «{nombre}» ({u.correo}) solo en «Reclutadora del centro» dentro del ambiente demo")

    # --- API: aislamiento, tablero, filtros, fichas y pipeline ---
    sp = db.begin_nested()  # sondas temporales; todo lo de aquí se deshace
    sondas = {}
    for clave, cuenta in cuentas.items():
        u = Usuario(correo=f"sonda.{secrets.token_hex(4)}@{DOMINIO_CORREO}", nombre=f"Sonda demo {cuenta.nombre}", rol="Administrador",
                    hash_pass="!", activo=True)
        db.add(u)
        db.flush()
        db.add(UsuarioCuenta(usuario_id=u.id, cuenta_id=cuenta.id))
        sondas[clave] = u
    db.flush()
    commit_real = db.commit
    db.commit = db.flush  # ningún endpoint puede confirmar la transacción de la carga
    actual = {"u": None}

    def _db():
        yield db

    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[usuario_actual] = lambda: actual["u"]
    cliente = TestClient(app)
    try:
        for clave, cuenta in cuentas.items():
            actual["u"] = sondas[clave]
            otra = next(c for k, c in cuentas.items() if k != clave)
            propias = por_cuenta[clave]
            ajenas = set(vac_excel) - propias
            print(f"\n— API como usuario de «{cuenta.nombre}»")
            tablero = cliente.get("/candidatos").json()
            demo_t = [t for t in tablero if t["candidatoCodigo"].startswith(PREFIJO_CAND)]
            val.check(all(t["vacanteId"] in propias for t in demo_t), "el tablero no muestra candidatos de la otra Cuenta")
            esperados_c = sum(n for vid in propias for n in esperado[vid].values())
            val.check(len(demo_t) == esperados_c, f"tablero: {len(demo_t)} tarjetas demo (Excel {esperados_c})")
            lista_v = {v["id"]: v for v in cliente.get("/vacantes").json()}
            val.check(set(lista_v) & ajenas == set() and propias <= set(lista_v), f"vacantes visibles = las {len(propias)} propias, ninguna ajena")
            muestra_ajena = sorted(ajenas)[0]
            val.check(cliente.get(f"/vacantes/{muestra_ajena}").status_code == 404, f"GET /vacantes/{muestra_ajena} (otra Cuenta) → 404")
            p_ajena = db.query(Postulacion).filter(Postulacion.cuenta_id == otra.id).first()
            if p_ajena:
                val.check(cliente.get(f"/candidatos/{p_ajena.codigo}").status_code == 404, f"GET /candidatos/{p_ajena.codigo} (otra Cuenta) → 404")
            val.check(cliente.get("/candidatos", headers={"X-Cuenta-Id": str(otra.id)}).json() == tablero,
                      "pedir la otra Cuenta por cabecera no cambia nada (sin vínculo, sin acceso)")

            malos_v, malos_f, malos_ficha = [], [], []
            for vid in sorted(propias):
                v = lista_v[vid]
                esp = esperado[vid]
                if v["candidatos"] != sum(esp.values()) or {e: n for e, n in (v["embudo"].get("etapas") or {}).items() if n} != esp:
                    malos_v.append(vid)
                for etapa in ETAPAS_CANDIDATO:
                    filtrado = cliente.get("/candidatos", params={"vacante": vid, "etapa": etapa}).json()
                    if len(filtrado) != esp.get(etapa, 0):
                        malos_f.append(f"{vid}/{etapa}: {len(filtrado)} vs {esp.get(etapa, 0)}")
                    if filtrado:
                        ficha = cliente.get(f"/candidatos/{filtrado[0]['id']}").json()
                        nivel = ORDEN[etapa]
                        ok = (ficha.get("etapa") == etapa and ficha.get("vacanteId") == vid
                              and (nivel < ORDEN["Entrevista IA"] or (ficha.get("entrevistaStatus") or {}).get("estado") == "evaluada")
                              and (nivel < ORDEN["Entrevista Humana"] or ficha.get("entrevistaHumana"))
                              and (nivel < ORDEN["Contratación"] or ficha.get("expedienteId")))
                        if not ok:
                            malos_ficha.append(f"{vid}/{etapa}")
            val.check(not malos_v, "tarjeta de cada vacante (total y embudo por etapa) = Excel" + (f": {malos_v}" if malos_v else ""))
            val.check(not malos_f, "filtro del tablero vacante × etapa = Excel en todas las combinaciones" + (f": {malos_f[:3]}" if malos_f else ""))
            val.check(not malos_ficha, "la ficha de cada combinación coincide con su tarjeta (etapa, vacante, expediente)" + (f": {malos_ficha[:3]}" if malos_ficha else ""))
            pipeline = cliente.get("/metricas/pipeline").json()["candidatos"]["por_etapa"]
            conteo_tablero = Counter(t["etapa"] for t in tablero)
            val.check(all(pipeline.get(e, 0) == conteo_tablero.get(e, 0) for e in ETAPAS_CANDIDATO),
                      "pipeline por etapa = tablero: " + ", ".join(f"{e} {pipeline.get(e, 0)}" for e in ETAPAS_CANDIDATO))
    finally:
        app.dependency_overrides.clear()
        db.commit = commit_real
        sp.rollback()  # las sondas desaparecen
    restos = db.query(Usuario).filter(Usuario.correo.like(f"sonda.%@{DOMINIO_CORREO}")).count()
    val.check(restos == 0, "los usuarios sonda de la validación no quedaron en la base")


# ============================================================
# Principal
# ============================================================


def main() -> int:
    ap = argparse.ArgumentParser(description="Carga el ambiente de demostración desde el Excel.")
    ap.add_argument("--ejecutar-demo", action="store_true", help="aplica la carga (sin esto es un simulacro que no escribe)")
    ap.add_argument("--excel", type=Path, default=EXCEL_DEFAULT)
    ap.add_argument("--correo", action="append", default=[], metavar="NOMBRE=CORREO", help="usuario registrado de un integrante")
    args = ap.parse_args()

    overrides = {}
    for par in args.correo:
        if "=" not in par:
            print(f"--correo inválido: {par}")
            return 2
        nombre, correo = par.split("=", 1)
        overrides[norm(nombre)] = correo.strip()

    modo = "EJECUCIÓN" if args.ejecutar_demo else "SIMULACRO (no escribe nada)"
    print(f"=== Carga del ambiente demo — {modo}")
    print(f"    Base: {engine.url.render_as_string(hide_password=True)}")
    print(f"    Excel: {args.excel}")
    bloqueadas = bloquear_comunicaciones()
    print(f"    Comunicaciones bloqueadas: {bloqueadas} funciones de envío (WhatsApp, correo, notificaciones, Teams)")

    try:
        x = leer_excel(args.excel)
        esperado = validar_excel(x)
    except AbortarCarga as ex:
        print(f"\n❌ {ex}")
        return 1
    print(f"    Excel consistente: {len(x['vacantes'])} vacantes, {len(x['candidatos'])} candidatos")

    faltan = []
    insp = sa_inspect(engine)
    tablas = set(insp.get_table_names())
    for t in Base.metadata.sorted_tables:
        if t.name in tablas:
            existentes = {c["name"] for c in insp.get_columns(t.name)}
            faltan += [f"{t.name}.{c.name}" for c in t.columns if c.name not in existentes]
    if faltan:
        print(f"\n❌ El esquema de la base está atrasado ({', '.join(faltan[:5])}…). Arranca la API con esta versión una vez y vuelve a correr.")
        return 1

    base = datetime.now(timezone.utc).replace(hour=15, minute=0, second=0, microsecond=0)
    db = SessionLocal()
    val = Validador()
    try:
        reales_antes = huella(db, solo_reales=True)
        mk, cuentas, admins = cargar(db, x, overrides, base)
        db.flush()

        print("\n— Carga")
        for tipo in sorted(set(mk.creados) | set(mk.actualizados)):
            print(f"    {tipo}: {mk.creados.get(tipo, 0)} creados · {mk.actualizados.get(tipo, 0)} actualizados")
        for aviso in mk.avisos:
            print(f"    ⚠️  {aviso}")

        # Segunda pasada completa: no debe crear ni cambiar NADA.
        print("\n— Repetición de la carga (idempotencia)")
        todo_1 = huella(db, solo_reales=False)
        mk2, _, _ = cargar(db, x, overrides, base)
        db.flush()
        todo_2 = huella(db, solo_reales=False)
        val.check(not mk2.creados, "repetir la carga no crea registros" + (f": {dict(mk2.creados)}" if mk2.creados else ""))
        distintas = [t for t in todo_1 if todo_1[t] != todo_2.get(t)]
        val.check(not distintas, "repetir la carga no modifica ninguna fila" + (f" (cambió: {distintas})" if distintas else ""))
        reales_despues = huella(db, solo_reales=True)
        tocadas = [t for t in reales_antes if reales_antes[t] != reales_despues.get(t)]
        val.check(not tocadas, "datos reales intactos (huella de todas las tablas fuera del ambiente demo)" + (f" — cambió: {tocadas}" if tocadas else ""))

        validar(db, x, esperado, cuentas, admins, val)
        print("\n— Comunicaciones")
        val.check(not INTENTOS_BLOQUEADOS, "ningún WhatsApp, correo, notificación ni publicación salió" + (f": {INTENTOS_BLOQUEADOS[:5]}" if INTENTOS_BLOQUEADOS else ""))
    except AbortarCarga as ex:
        db.rollback()
        db.close()
        print(f"\n❌ {ex}\n   Nada se escribió (rollback).")
        return 1
    except Exception:
        db.rollback()
        db.close()
        print("\n❌ Error inesperado; nada se escribió (rollback).")
        raise

    if val.fallas:
        db.rollback()
        db.close()
        print(f"\n❌ {len(val.fallas)} validación(es) fallaron; nada se escribió (rollback).")
        return 1
    if not args.ejecutar_demo:
        db.rollback()
        db.close()
        print(f"\n🟡 SIMULACRO OK: {val.ok} comprobaciones pasaron. No se escribió nada.")
        print("   Para aplicar: .venv/Scripts/python.exe scripts/cargar_ambiente_demo.py --ejecutar-demo")
        return 0

    registrar(db, ACTOR, "carga_ambiente_demo", "cuenta", ",".join(str(c.id) for c in cuentas.values()), {
        "excel": args.excel.name, "creados": dict(mk.creados), "actualizados": dict(mk.actualizados),
        "avisos": mk.avisos, "comprobaciones": val.ok,
    })
    db.commit()
    db.close()
    print(f"\n🎉 Ambiente demo cargado: {val.ok} comprobaciones OK. Cuentas: "
          + ", ".join(f"{c.nombre} (id {c.id})" for c in cuentas.values()))
    return 0


if __name__ == "__main__":
    sys.exit(main())
